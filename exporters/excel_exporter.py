import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from models.creator import Creator
from utils.logger import get_logger

logger = get_logger(__name__)

class ExcelExporter:
    """Excel formatında dışa aktarma işlemlerini yöneten sınıf."""

    def __init__(self):
        self.header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
        self.header_font = Font(color="FFFFFF", bold=True)
        self.high_score_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        self.med_score_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        self.low_score_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        self.alt_row_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    def export(self, creators: list[Creator], keyword: str, output_dir: str = "data/results") -> str:
        """
        İçerik üreticilerini Excel formatında dışa aktarır.
        """
        logger.info(f"'{keyword}' için Excel raporu oluşturuluyor...")
        
        if not os.path.exists(output_dir):
            os.makedirs(output_dir, exist_ok=True)
            logger.debug(f"Çıktı dizini oluşturuldu: {output_dir}")

        if keyword.endswith(".xlsx"):
            filename = keyword
        else:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_keyword = keyword.replace(" ", "_").lower()
            filename = f"{safe_keyword}_{timestamp}.xlsx"
            
        filepath = os.path.join(output_dir, filename)

        wb = Workbook()
        
        # 1. Özet Sayfası
        ws_summary = wb.active
        ws_summary.title = "Özet"
        self._create_summary_sheet(ws_summary, creators, keyword)

        # Tüm Sonuçlar için ortak sütunlar
        columns = [
            "Sıra", "Kullanıcı Adı", "Platform", "Takipçi", "Etkileşim Oranı (%)", 
            "Skor", "Konu Etiketleri", "Bio", "Profil URL",
            "LLM Özeti", "Niş Alanı", "Hedef Kitle"
        ]

        # 2. Tüm Sonuçlar Sayfası
        ws_all = wb.create_sheet("Tüm Sonuçlar")
        sorted_creators = sorted(creators, key=lambda c: c.score if hasattr(c, 'score') and c.score else 0, reverse=True)
        self._populate_creators_sheet(ws_all, sorted_creators, columns)

        # 3. Platformlara Özel Sayfalar
        platforms = {"YouTube", "TikTok", "Instagram"}
        for platform in platforms:
            platform_creators = [c for c in sorted_creators if hasattr(c, 'platform') and c.platform and c.platform.lower() == platform.lower()]
            if platform_creators:
                ws_platform = wb.create_sheet(platform)
                self._populate_creators_sheet(ws_platform, platform_creators, columns)

        # Sütun genişliklerini ayarla
        for sheet in wb.sheetnames:
            self._adjust_column_widths(wb[sheet])

        wb.save(filepath)
        logger.info(f"Excel raporu başarıyla oluşturuldu: {filepath}")
        return filepath

    def _create_summary_sheet(self, ws, creators: list[Creator], keyword: str):
        ws.append(["Rapor Özeti", ""])
        ws["A1"].font = Font(bold=True, size=14)
        
        total_creators = len(creators)
        
        scores = [c.score for c in creators if hasattr(c, 'score') and c.score is not None]
        avg_score = sum(scores) / len(scores) if scores else 0
        
        engagements = [c.engagement_rate for c in creators if hasattr(c, 'engagement_rate') and c.engagement_rate is not None]
        avg_eng = sum(engagements) / len(engagements) if engagements else 0
        
        platforms = {}
        for c in creators:
            if hasattr(c, 'platform') and c.platform:
                platforms[c.platform] = platforms.get(c.platform, 0) + 1

        data = [
            ["Anahtar Kelime:", keyword],
            ["Toplam İçerik Üreticisi:", total_creators],
            ["Ortalama Skor:", round(avg_score, 2)],
            ["Ortalama Etkileşim (%):", round(avg_eng, 2)],
            ["", ""],
            ["Platform Dağılımı", ""]
        ]
        
        for p, count in platforms.items():
            data.append([p, count])
            
        for row in data:
            ws.append(row)

    def _populate_creators_sheet(self, ws, creators: list[Creator], columns: list[str]):
        ws.append(columns)
        
        # Header formatting
        for col_idx, _ in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        ws.freeze_panes = "A2"

        for idx, c in enumerate(creators, 1):
            score = getattr(c, 'score', 0) or 0
            tags = ", ".join(getattr(c, 'tags', [])) if getattr(c, 'tags', []) else ""
            
            row_data = [
                idx,
                getattr(c, 'username', ''),
                getattr(c, 'platform', ''),
                getattr(c, 'followers', 0),
                getattr(c, 'engagement_rate', 0.0),
                score,
                tags,
                getattr(c, 'bio', ''),
                getattr(c, 'url', ''),
                getattr(c, 'llm_summary', ''),
                getattr(c, 'niche', ''),
                getattr(c, 'target_audience', '')
            ]
            
            ws.append(row_data)
            row_idx = idx + 1
            
            # Row formatting
            for col_idx in range(1, len(columns) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                
                # Alternating colors
                if row_idx % 2 == 0:
                    cell.fill = self.alt_row_fill
                    
                # Conditional formatting for score (column 6)
                if col_idx == 6:
                    if score > 80:
                        cell.fill = self.high_score_fill
                    elif score >= 50:
                        cell.fill = self.med_score_fill
                    else:
                        cell.fill = self.low_score_fill

    def _adjust_column_widths(self, ws):
        for col in ws.columns:
            max_length = 0
            column_letter = col[0].column_letter
            for cell in col:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
